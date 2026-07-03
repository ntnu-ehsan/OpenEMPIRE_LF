"""Build a PowerPoint deck from the figures produced by scenarios_comparison.ipynb.

Reads the PNG figures in Results/comparison_plots/ (exported by the notebook via
kaleido) and assembles a titled slide per figure with a short caption. Scenario
names, objective values and periods are pulled from the executed notebook so the
deck always matches the latest run.

Run:  python scripts/make_comparison_presentation.py
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

try:
    from PIL import Image
    _HAVE_PIL = True
except Exception:
    _HAVE_PIL = False

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parent.parent
PLOT_DIR = ROOT / "Results" / "comparison_plots"
NOTEBOOK = ROOT / "scripts" / "scenarios_comparison.ipynb"
OUT_PPTX = ROOT / "Results" / "scenario_comparison.pptx"

# ── Palette ───────────────────────────────────────────────────────────────────
DARK = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x1B, 0x8A, 0x8F)
ACCENT2 = RGBColor(0xE8, 0x7A, 0x2B)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
GREY = RGBColor(0x55, 0x5B, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ── Slide order: (figure stem, title, caption) ───────────────────────────────
# Region/North-Sea variants follow their Europe-wide parent.
SLIDES = [
    ("01_system_cost", "System Cost (Objective Value)",
     "Total discounted system cost per scenario over the shared horizon."),
    ("02_installed_capacity_stacked", "Installed Generation Capacity by Technology",
     "Europe-wide installed capacity by technology; 'Initial' = existing fleet. "
     "Original (solid) vs aggregated (hatched) runs paired by period."),
    ("02b_installed_capacity_region", "Installed Capacity — North Sea region",
     "Same metric restricted to the studied North Sea nodes (DE, DK, NL, NO, UK "
     "at NUTS-2, plus Belgium and offshore)."),
    ("03_capacity_by_tech_lines", "Capacity Trajectory per Technology",
     "One panel per technology group; each line is a scenario from the initial "
     "fleet through the horizon."),
    ("03b_capacity_by_tech_lines_region", "Capacity Trajectory — North Sea region",
     "Per-technology capacity trajectories summed over the North Sea nodes."),
    ("04_annual_generation_stacked", "Annual Electricity Generation",
     "Europe-wide expected annual generation by technology."),
    ("04b_annual_generation_region", "Annual Generation — North Sea region",
     "Expected annual generation by technology over the North Sea nodes."),
    ("05_renewable_share", "Renewable Generation Share",
     "Share of renewable generation over time per scenario."),
    ("06_co2_emissions_price", "CO₂ Emissions & Carbon Price",
     "Total CO₂ emissions and the resulting CO₂ price per period."),
    ("07_co2_intensity", "Average CO₂ Emission Factor",
     "System-average carbon intensity of electricity over time."),
    ("08_electricity_price", "Average Electricity Price",
     "Europe-wide average electricity price per period and scenario."),
    ("08b_electricity_price_region", "Electricity Price — North Sea region",
     "Demand-weighted average price across the North Sea nodes."),
    ("11_nodal_price_maps_region", "Nodal Price Maps — North Sea",
     "Demand-weighted average nodal price (first vs last period) per scenario. "
     "Nationally-resolved countries are imputed from REPowerEU++ NUTS-2 detail, "
     "rescaled to their national price."),
    ("09_curtailment_losses", "Curtailed RES & System Losses",
     "Curtailed renewable energy and transmission/storage losses per scenario."),
    ("10_storage_capacity", "Storage Power & Energy Capacity",
     "Total installed storage power and energy capacity over time."),
    ("11_storage_by_type", "Storage by Technology Type",
     "Installed storage power capacity split by storage technology."),
    ("12_load_shedding", "Load Shedding",
     "Expected annual load shedding per period and scenario."),
    ("12b_load_shedding_region", "Load Shedding — North Sea region",
     "Approximate annual load shedding summed over the North Sea nodes."),
    ("13_capacity_delta", "Capacity Delta — First vs Last Period",
     "Net change in installed capacity by technology between the first and last "
     "common period."),
]


# ── Notebook metadata extraction ──────────────────────────────────────────────
def notebook_streams() -> str:
    if not NOTEBOOK.exists():
        return ""
    nb = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    out = []
    for c in nb.get("cells", []):
        for o in c.get("outputs", []):
            if o.get("output_type") == "stream":
                out.append("".join(o.get("text", [])))
    return "\n".join(out)


def parse_run_info(text: str) -> dict:
    info = {"scenarios": [], "periods": [], "objectives": {}}
    for m in re.finditer(r"^\s{2}(\S.*?)\s{2,}periods=.*?obj=([\d.]+|N/A)", text, re.M):
        name = m.group(1).strip()
        info["scenarios"].append(name)
        if m.group(2) != "N/A":
            info["objectives"][name] = float(m.group(2))
    pm = re.search(r"Common periods \(\d+\):\s*\[(.*?)\]", text)
    if pm:
        info["periods"] = re.findall(r"'([^']+)'", pm.group(1))
    return info


# ── Slide helpers ─────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, text, size=18, color=DARK, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def picture_fit(slide, img_path: Path, box_x, box_y, box_w, box_h):
    """Add a picture scaled to fit the box, preserving aspect ratio, centered."""
    if _HAVE_PIL:
        with Image.open(img_path) as im:
            iw, ih = im.size
        ar = iw / ih
        box_ar = box_w / box_h
        if ar >= box_ar:           # width-limited
            w = box_w
            h = int(box_w / ar)
        else:                      # height-limited
            h = box_h
            w = int(box_h * ar)
        x = box_x + (box_w - w) // 2
        y = box_y + (box_h - h) // 2
        slide.shapes.add_picture(str(img_path), x, y, width=Emu(w), height=Emu(h))
    else:
        slide.shapes.add_picture(str(img_path), box_x, box_y, width=box_w)


# ── Title slide ───────────────────────────────────────────────────────────────
def build_title(info: dict):
    s = add_slide()
    rect(s, 0, 0, SW, SH, DARK)
    rect(s, 0, Inches(2.55), SW, Inches(0.06), ACCENT2)
    textbox(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(1.0),
            "OpenEMPIRE — North Sea Scenario Comparison",
            size=40, color=WHITE, bold=True)
    textbox(s, Inches(0.9), Inches(2.75), Inches(11.5), Inches(0.7),
            "Capacity, generation, prices, emissions and storage across scenarios",
            size=20, color=LIGHT)
    scen = ", ".join(info["scenarios"]) or "see notebook"
    per = info["periods"]
    span = f"{per[0]} → {per[-1]}" if per else "common horizon"
    textbox(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.6),
            f"Scenarios:  {scen}\nHorizon:  {span}  ({len(per)} periods)\n"
            "Source:  scripts/scenarios_comparison.ipynb",
            size=18, color=LIGHT)


# ── Summary slide (objectives) ────────────────────────────────────────────────
def build_summary(info: dict):
    objs = info.get("objectives", {})
    if not objs:
        return
    s = add_slide()
    rect(s, 0, 0, SW, Inches(1.0), DARK)
    textbox(s, Inches(0.5), Inches(0.15), Inches(12), Inches(0.7),
            "Total Discounted System Cost", size=26, color=WHITE, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    y = Inches(1.6)
    lo = min(objs.values())
    for name, val in objs.items():
        rect(s, Inches(0.8), y, Inches(3.2), Inches(0.55), LIGHT)
        textbox(s, Inches(1.0), y, Inches(3.0), Inches(0.55), name,
                size=16, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        bar_full = Inches(7.5)
        w = int(bar_full * (val / max(objs.values())))
        rect(s, Inches(4.2), y, Emu(w), Inches(0.55),
             ACCENT if val == lo else ACCENT2)
        textbox(s, Inches(4.3) + Emu(w), y, Inches(2.2), Inches(0.55),
                f"{val:,.0f} B€", size=15, color=GREY,
                anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.8)
    textbox(s, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.6),
            "Lowest-cost scenario highlighted in teal.", size=13, color=GREY)


# ── Figure slides ─────────────────────────────────────────────────────────────
def build_figure(stem: str, title: str, caption: str) -> bool:
    img = PLOT_DIR / f"{stem}.png"
    if not img.exists():
        print(f"  skip (no PNG): {stem}")
        return False
    s = add_slide()
    rect(s, 0, 0, SW, Inches(0.95), DARK)
    textbox(s, Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.7),
            title, size=24, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    picture_fit(s, img, Inches(0.4), Inches(1.1), Inches(12.55), Inches(5.25))
    rect(s, 0, Inches(6.55), SW, Inches(0.95), LIGHT)
    textbox(s, Inches(0.5), Inches(6.62), Inches(12.3), Inches(0.85),
            caption, size=13, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    return True


def main():
    info = parse_run_info(notebook_streams())
    print(f"Scenarios: {info['scenarios']}  periods: {len(info['periods'])}")
    build_title(info)
    build_summary(info)
    n = 0
    for stem, title, caption in SLIDES:
        if build_figure(stem, title, caption):
            n += 1
    prs.save(OUT_PPTX)
    print(f"Wrote {OUT_PPTX}  ({n} figure slides, {len(prs.slides)} total)")


if __name__ == "__main__":
    main()
