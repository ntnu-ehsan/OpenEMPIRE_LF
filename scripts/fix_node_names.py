"""
Fix node name mismatches in north_sea dataset Excel files.

The Nodes sheet uses names with spaces (e.g., "Dogger Bank") but
BidirectionalLines, DirectionalLines, and Transmission.xlsx use 
names without spaces (e.g., "DoggerBank"). This script normalizes
all names to match the Nodes sheet.

Usage:
    python scripts/fix_node_names.py
"""

import argparse
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

parser = argparse.ArgumentParser(description="Fix node name mismatches in a north_sea dataset.")
parser.add_argument("dataset", nargs="?", default="north_sea",
                    help="Dataset folder name under 'Data handler/' (default: north_sea)")
args = parser.parse_args()
DATA_DIR = Path("Data handler") / args.dataset

# Mapping: wrong name -> correct name (matching Nodes sheet)
NAME_FIXES = {
    "DoggerBank": "Dogger Bank",
    "EastAnglia": "East Anglia",
    "FirthofForth": "Firth of Forth",
    "HelgolanderBucht": "Helgolander Bucht",
    "HollandseeKust": "Hollandsee Kust",
    "MorayFirth": "Moray Firth",
    "OuterDowsing": "Outer Dowsing",
    "SorligeNordsjoI": "Sorlige Nordsjo I",
    "SorligeNordsjoII": "Sorlige Nordsjo II",
    "UtsiraNord": "Utsira Nord",
    "BosniaH": "Bosnia H",
    "CzechR": "Czech R",
    "GreatBrit.": "Great Brit.",
}


def fix_workbook(filepath: Path):
    """Fix node names in all sheets of an Excel workbook."""
    wb = load_workbook(filepath)
    changes = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    stripped = cell.value.strip()
                    if stripped in NAME_FIXES:
                        cell.value = NAME_FIXES[stripped]
                        changes += 1
    if changes > 0:
        wb.save(filepath)
        print(f"  {filepath.name}: {changes} cells fixed")
    else:
        print(f"  {filepath.name}: no changes needed")
    return changes


def main():
    files_to_check = [
        DATA_DIR / "Sets.xlsx",
        DATA_DIR / "Transmission.xlsx",
        DATA_DIR / "Generator.xlsx",
        DATA_DIR / "General.xlsx",
    ]

    # Also check Node.xlsx if it exists
    node_file = DATA_DIR / "Node.xlsx"
    if node_file.exists():
        files_to_check.append(node_file)

    total = 0
    for f in files_to_check:
        if f.exists():
            total += fix_workbook(f)
        else:
            print(f"  {f.name}: not found, skipping")

    print(f"\nTotal cells fixed: {total}")
    if total > 0:
        print("Done! Re-run the model after pushing these changes.")
    else:
        print("No mismatches found.")


if __name__ == "__main__":
    main()
